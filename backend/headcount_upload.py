"""Upload and load weekly personnel counts stored in the SharePoint workbook."""

from __future__ import annotations

import base64
import io
import re
import time
import unicodedata
from typing import Any

import openpyxl
import requests
import streamlit as st

from data_extractor import SHAREPOINT_URL_HEADCOUNT


SHEET_MARKER = "CONTEO DE PERSONAL SEMANAL"
DASHBOARD_RANCHES = {
    "ramona": "Ramona",
    "isabela": "Isabela",
    "christina": "Christina",
    "cecilia": "Cecilia",
    "poscosecha": "Poscosecha",
    "vivero": "Propagacion",
}
DASHBOARD_GENERAL_SECTIONS = {
    "operativos": "Operativo",
}
DASHBOARD_CONCEPTS = {
    "manejo de planta": "Manejo P.",
    "trasplante": "Trasplante",
    "hoops": "Hoops",
    "supervisor": "Supervisores",
    "siembra": "Siembra",
    "consolidador": "Consolidacion",
    "riego": "Riego",
    "movimiento de charola": "Mov. Charolas",
    "secado st": "Esquejes",
    "llenado de charola": "Esquejes",
    "maquina": "Siembra",
    "mipe": "MIPE Y MIRFE",
    "mirfe": "MIPE Y MIRFE",
    "corte": "Corte",
}
POSCOSECHA_CONCEPTS = {
    "upc": "Alm.upc y empaq",
    "bouquetera": "Prod. Patina y rec",
    "supervisor": "Supervisores",
    "recepcion": "Prod. Patina y rec",
    "almacen": "Alm.upc y empaq",
    "patinador": "Prod. Patina y rec",
    "empaque": "Alm.upc y empaq",
    "consumer": "Prod. Patina y rec",
    "cortador": "Prod. Patina y rec",
    "aux. limpieza": "Alm.upc y empaq",
}
OPERATIVO_CONCEPTS = {
    "chofer": "Transporte",
    "cameros": "Tract. Y Cameros",
    "tractores": "Tract. Y Cameros",
    "hoops": "Hoops",
    "soldador": "Soldadores",
}


def _normalize(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    return " ".join("".join(ch for ch in text if not unicodedata.combining(ch)).casefold().split())


def _as_count(value: Any, *, row_number: int, column_name: str) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        raise ValueError(f"El valor de {column_name} en la fila {row_number} no es un conteo numérico.")
    if not number.is_integer() or number < 0:
        raise ValueError(f"El valor de {column_name} en la fila {row_number} debe ser un entero no negativo.")
    return int(number)


def _parse_week_code(week_code: str) -> tuple[str, int, int]:
    code = str(week_code or "").strip()
    if not re.fullmatch(r"\d{4}", code):
        raise ValueError("El código de semana debe tener cuatro dígitos, por ejemplo 2638.")
    year, week = 2000 + int(code[:2]), int(code[2:])
    if not 1 <= week <= 53:
        raise ValueError("El número de semana debe estar entre 01 y 53.")
    return code, year, week


def parse_personnel_workbook(content: bytes) -> list[dict[str, Any]]:
    """Read the Planta and Contratistas count columns; dollar columns are ignored."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo abrir el archivo. Selecciona el Excel semanal .xlsx.") from exc

    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, ())
    subheader = next(rows, ())
    if (
        _normalize(header[0] if len(header) > 0 else None) != "area"
        or _normalize(subheader[1] if len(subheader) > 1 else None) != "planta"
        or _normalize(subheader[2] if len(subheader) > 2 else None) != "contratistas"
        or _normalize(subheader[3] if len(subheader) > 3 else None) != "total"
    ):
        raise ValueError("El formato no coincide con AF.xlsx: se esperan las columnas Área, Planta, Contratistas y Total.")

    result: list[dict[str, Any]] = []
    current_section = ""
    for row_number, row in enumerate(rows, start=3):
        padded = list(row) + [None] * max(0, 4 - len(row))
        label = str(padded[0] or "").strip()
        if not label:
            continue
        if _normalize(label) in {"area", "totales", "total"}:
            continue

        plant = _as_count(padded[1], row_number=row_number, column_name="Planta")
        contractors = _as_count(padded[2], row_number=row_number, column_name="Contratistas")
        supplied_total = _as_count(padded[3], row_number=row_number, column_name="Total")

        if plant is None and contractors is None and supplied_total is None:
            current_section = label
            continue
        if plant is None and contractors is None:
            raise ValueError(f"Falta separar el conteo de Planta y Contratistas en la fila {row_number}.")

        plant = plant or 0
        contractors = contractors or 0
        total = plant + contractors
        if supplied_total is not None and supplied_total != total:
            raise ValueError(f"El total de la fila {row_number} no coincide con Planta + Contratistas.")
        if not current_section:
            raise ValueError(f"No se identificó el rancho o sección antes de la fila {row_number}.")

        result.append({
            "source_section": current_section,
            "concept": label,
            "planta": plant,
            "contratistas": contractors,
            "total": total,
        })

    workbook.close()
    if not result:
        raise ValueError("El archivo no contiene filas de conteo para importar.")
    return result


def _graph_credentials() -> tuple[dict[str, str], str]:
    if not SHAREPOINT_URL_HEADCOUNT:
        raise RuntimeError("Falta configurar el libro de conteo de SharePoint en el servidor.")
    try:
        credentials = st.secrets["sharepoint"]
        tenant_id = credentials["tenant_id"]
        client_id = credentials["client_id"]
        client_secret = credentials["client_secret"]
    except (KeyError, TypeError) as exc:
        raise RuntimeError("Falta la configuración de acceso a SharePoint en el servidor.") from exc

    token_response = requests.post(
        f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
        },
        timeout=20,
    )
    if token_response.status_code != 200:
        raise RuntimeError("No se pudo autenticar el servidor con SharePoint.")
    token = token_response.json().get("access_token")
    if not token:
        raise RuntimeError("SharePoint no devolvió un token de acceso válido.")
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}, token


def _resolve_workbook(headers: dict[str, str]) -> tuple[str, str]:
    share_token = base64.urlsafe_b64encode(SHAREPOINT_URL_HEADCOUNT.encode("utf-8")).decode("ascii").rstrip("=")
    response = requests.get(
        f"https://graph.microsoft.com/v1.0/shares/u!{share_token}/driveItem",
        headers=headers,
        timeout=20,
    )
    if response.status_code != 200:
        raise RuntimeError(f"No se pudo abrir el libro de SharePoint (HTTP {response.status_code}).")
    item = response.json()
    drive_id = item.get("parentReference", {}).get("driveId")
    item_id = item.get("id")
    if not drive_id or not item_id:
        raise RuntimeError("SharePoint no devolvió la referencia del libro.")
    return drive_id, item_id


def _workbook_session() -> tuple[str, dict[str, str]]:
    headers, _ = _graph_credentials()
    drive_id, item_id = _resolve_workbook(headers)
    workbook_url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/workbook"
    response = requests.post(
        f"{workbook_url}/createSession",
        headers=headers,
        json={"persistChanges": True},
        timeout=30,
    )
    if response.status_code not in (200, 201):
        raise RuntimeError(f"No se pudo abrir una sesión para guardar en SharePoint (HTTP {response.status_code}).")
    session_id = response.json().get("id")
    if not session_id:
        raise RuntimeError("SharePoint no devolvió una sesión de edición.")
    return workbook_url, {**headers, "workbook-session-id": session_id}


def save_personnel_week(week_code: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    code, _, _ = _parse_week_code(week_code)
    workbook_url, headers = _workbook_session()
    worksheets_url = f"{workbook_url}/worksheets"
    sheets_response = requests.get(worksheets_url, headers=headers, timeout=20)
    if sheets_response.status_code != 200:
        raise RuntimeError("No se pudieron consultar las pestañas del libro de SharePoint.")
    existing_names = {
        str(sheet.get("name", "")).casefold()
        for sheet in sheets_response.json().get("value", [])
    }

    if code.casefold() in existing_names:
        marker_response = requests.get(
            f"{worksheets_url}/{code}/range(address='A1')?$select=values",
            headers=headers,
            timeout=20,
        )
        if marker_response.status_code != 200:
            raise RuntimeError(f"No se pudo revisar la pestaña {code} antes de actualizarla.")
        marker_values = marker_response.json().get("values", [[None]])
        marker = marker_values[0][0] if marker_values and marker_values[0] else None
        if marker not in (None, "", SHEET_MARKER):
            raise RuntimeError(f"La pestaña {code} ya existe y contiene otros datos; no se modificó.")
        clear_response = requests.post(
            f"{worksheets_url}/{code}/range(address='A1:E1000')/clear",
            headers=headers,
            json={"applyTo": "Contents"},
            timeout=20,
        )
        if clear_response.status_code not in (200, 204):
            raise RuntimeError(f"No se pudo preparar la pestaña {code} para reemplazar el conteo.")
    else:
        add_response = requests.post(
            f"{worksheets_url}/add",
            headers=headers,
            json={"name": code},
            timeout=20,
        )
        if add_response.status_code not in (200, 201):
            raise RuntimeError(f"No se pudo crear la pestaña semanal {code} en SharePoint.")
        time.sleep(0.5)

    matrix = [
        [SHEET_MARKER, code, "", "", ""],
        ["Rancho/sección (Excel)", "Concepto", "Planta", "Contratistas", "Total"],
    ]
    matrix.extend([
        [row["source_section"], row["concept"], row["planta"], row["contratistas"], row["total"]]
        for row in rows
    ])
    end_row = len(matrix)
    write_response = requests.patch(
        f"{worksheets_url}/{code}/range(address='A1:E{end_row}')",
        headers=headers,
        json={"values": matrix},
        timeout=60,
    )
    if write_response.status_code not in (200, 201):
        raise RuntimeError(f"No se pudo escribir el conteo en la pestaña {code} (HTTP {write_response.status_code}).")

    try:
        requests.patch(
            f"{worksheets_url}/{code}/range(address='A2:E2')/format/fill",
            headers=headers,
            json={"color": "#0F766E"},
            timeout=10,
        )
        requests.patch(
            f"{worksheets_url}/{code}/range(address='A2:E2')/format/font",
            headers=headers,
            json={"color": "#FFFFFF", "bold": True},
            timeout=10,
        )
        requests.patch(
            f"{worksheets_url}/{code}/range(address='C3:E{end_row}')/format",
            headers=headers,
            json={"numberFormat": "0"},
            timeout=10,
        )
    except requests.RequestException:
        pass

    known_sections = {
        _normalize(row["source_section"])
        for row in rows
        if _normalize(row["source_section"]) in DASHBOARD_RANCHES
        or _normalize(row["source_section"]) in DASHBOARD_GENERAL_SECTIONS
    }
    all_sections = {_normalize(row["source_section"]): row["source_section"] for row in rows}
    mapped_sections = sorted(
        [DASHBOARD_RANCHES[key] for key in known_sections if key in DASHBOARD_RANCHES]
        + [DASHBOARD_GENERAL_SECTIONS[key] for key in known_sections if key in DASHBOARD_GENERAL_SECTIONS]
    )
    unmapped_sections = sorted(
        label for key, label in all_sections.items()
        if key not in DASHBOARD_RANCHES and key not in DASHBOARD_GENERAL_SECTIONS
    )
    return {
        "week_code": code,
        "sheet_name": code,
        "stored_rows": len(rows),
        "mapped_sections": mapped_sections,
        "unmapped_sections": unmapped_sections,
    }


def load_dashboard_headcounts() -> list[dict[str, Any]]:
    """Load only reviewed ranch/concept mappings from weekly sheets in SharePoint."""
    headers, _ = _graph_credentials()
    share_token = base64.urlsafe_b64encode(SHAREPOINT_URL_HEADCOUNT.encode("utf-8")).decode("ascii").rstrip("=")
    response = requests.get(
        f"https://graph.microsoft.com/v1.0/shares/u!{share_token}/driveItem/content",
        headers=headers,
        timeout=45,
    )
    if response.status_code != 200:
        raise RuntimeError(f"No se pudo leer el libro de conteos de SharePoint (HTTP {response.status_code}).")

    workbook = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    aggregated: dict[tuple[int, int, str], dict[str, Any]] = {}

    def get_record(subcat: str) -> dict[str, Any]:
        key = (year, week, subcat)
        return aggregated.setdefault(key, {
            "semana": int(code),
            "year": year,
            "week": week,
            "date_range": "",
            "subcat": subcat,
            "mxn_total": 0,
            "usd_total": 0,
            "mxn_ranches": {},
            "usd_ranches": {},
            "hc_planta_total": 0,
            "hc_contratistas_total": 0,
            "hc_planta_ranches": {},
            "hc_contratistas_ranches": {},
            "hc_operativo_total": 0,
            "hc_operativo_planta": 0,
            "hc_operativo_contratistas": 0,
        })

    def add_count(subcat: str, ranch: str, plant: int, contractors: int) -> None:
        record = get_record(subcat)
        record["hc_planta_total"] += plant
        record["hc_contratistas_total"] += contractors
        record["hc_planta_ranches"][ranch] = record["hc_planta_ranches"].get(ranch, 0) + plant
        record["hc_contratistas_ranches"][ranch] = record["hc_contratistas_ranches"].get(ranch, 0) + contractors

    def add_operativo_count(subcat: str, plant: int, contractors: int) -> None:
        record = get_record(subcat)
        record["hc_planta_total"] += plant
        record["hc_contratistas_total"] += contractors
        record["hc_operativo_total"] += plant + contractors
        record["hc_operativo_planta"] += plant
        record["hc_operativo_contratistas"] += contractors

    for worksheet in workbook.worksheets:
        if not re.fullmatch(r"\d{4}", worksheet.title):
            continue
        code, year, week = _parse_week_code(worksheet.title)
        marker_row = next(worksheet.iter_rows(min_row=1, max_row=1, max_col=1, values_only=True), (None,))
        if not marker_row or marker_row[0] != SHEET_MARKER:
            continue

        for row in worksheet.iter_rows(min_row=3, min_col=1, max_col=5, values_only=True):
            source_section, concept, plant_value, contractor_value, _ = row
            section_key = _normalize(source_section)
            concept_key = _normalize(concept)
            if section_key in DASHBOARD_GENERAL_SECTIONS:
                subcat = OPERATIVO_CONCEPTS.get(concept_key)
                if not subcat:
                    continue
                plant = _as_count(plant_value, row_number=0, column_name="Planta") or 0
                contractors = _as_count(contractor_value, row_number=0, column_name="Contratistas") or 0
                if plant or contractors:
                    add_operativo_count(subcat, plant, contractors)
                continue

            ranch = DASHBOARD_RANCHES.get(section_key)
            subcat = (
                POSCOSECHA_CONCEPTS.get(concept_key)
                if section_key == "poscosecha"
                else DASHBOARD_CONCEPTS.get(concept_key)
            )
            if not ranch or not subcat:
                continue
            plant = _as_count(plant_value, row_number=0, column_name="Planta") or 0
            contractors = _as_count(contractor_value, row_number=0, column_name="Contratistas") or 0
            if ranch == "Propagacion" and _normalize(concept) == "movimiento de charola":
                contractors_for_consolidation = contractors // 2
                contractors_for_planting = contractors - contractors_for_consolidation
                add_count(subcat, ranch, plant, 0)
                add_count("Consolidacion", ranch, 0, contractors_for_consolidation)
                add_count("Siembra", ranch, 0, contractors_for_planting)
                continue

            add_count(subcat, ranch, plant, contractors)

    workbook.close()
    return list(aggregated.values())


def add_dashboard_headcounts(data: dict[str, Any]) -> None:
    records = load_dashboard_headcounts()
    if not records:
        return
    data.setdefault("mano_obra_data", []).extend(records)
    years = set(data.get("years", []))
    weeks_per_year = data.setdefault("weeks_per_year", {})
    for record in records:
        years.add(record["year"])
        year_key = str(record["year"])
        stored_weeks = list(weeks_per_year.get(year_key, weeks_per_year.get(record["year"], [])) or [])
        if record["week"] not in stored_weeks:
            stored_weeks.append(record["week"])
        weeks_per_year[year_key] = sorted(stored_weeks)
    data["years"] = sorted(years)
